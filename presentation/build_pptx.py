#!/usr/bin/env python3
"""Generate КРУЖИМ Brand Positioning 2026 PowerPoint presentation."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import nsmap
from lxml import etree

# Brand colors
INK = RGBColor(0x0A, 0x1C, 0x24)
INK_MID = RGBColor(0x12, 0x30, 0x40)
TEAL = RGBColor(0x1A, 0x7A, 0x7A)
TEAL_DEEP = RGBColor(0x0E, 0x5C, 0x5C)
SIGNAL = RGBColor(0xE8, 0x5A, 0x2A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MIST = RGBColor(0xF3, 0xF7, 0xF9)
MUTED = RGBColor(0x5A, 0x71, 0x7C)
FOG = RGBColor(0xE7, 0xEE, 0xF2)
SOFT = RGBColor(0x8A, 0xA0, 0xAB)
DARK_BG = RGBColor(0x0D, 0x25, 0x30)


def set_run(run, size=18, bold=False, color=INK, font="Arial"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_textbox(slide, left, top, width, height, text, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT, font="Arial"):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, color=color, font=font)
    return box


def add_paragraph(tf, text, size=16, bold=False, color=INK, space_before=6, space_after=4, font="Arial"):
    p = tf.add_paragraph()
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, color=color, font=font)
    return p


def fill_shape(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def add_rect(slide, left, top, width, height, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    fill_shape(shape, color)
    return shape


def add_round_rect(slide, left, top, width, height, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    fill_shape(shape, color)
    # softer corners
    try:
        shape.adjustments[0] = 0.1
    except Exception:
        pass
    return shape


def footer(slide, page, total=9, dark=False):
    color = SOFT if not dark else RGBColor(0x6A, 0x82, 0x8C)
    add_textbox(slide, Inches(0.5), Inches(7.05), Inches(6), Inches(0.3),
                "КРУЖИМ · Brand book", size=11, color=color)
    add_textbox(slide, Inches(11.5), Inches(7.05), Inches(1.3), Inches(0.3),
                f"{page:02d} / {total:02d}", size=11, color=color, align=PP_ALIGN.RIGHT)


def topbar(slide, section, dark=False):
    brand_c = WHITE if dark else INK
    sec_c = RGBColor(0x9A, 0xB4, 0xBE) if dark else SOFT
    add_textbox(slide, Inches(0.5), Inches(0.28), Inches(3), Inches(0.35),
                "КРУЖИМ", size=14, bold=True, color=brand_c)
    add_textbox(slide, Inches(7.5), Inches(0.3), Inches(5.3), Inches(0.35),
                section.upper(), size=11, bold=True, color=sec_c, align=PP_ALIGN.RIGHT)


def blank_slide(prs, dark=False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = DARK_BG if dark else MIST
    fill_shape(slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height), bg)
    return slide


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # --- 1 Title ---
    s = blank_slide(prs, dark=True)
    add_rect(s, Inches(9.5), Inches(0), Inches(3.9), Inches(7.5), RGBColor(0x12, 0x33, 0x3D))
    topbar(s, "Brand Positioning · 2026", dark=True)
    add_textbox(s, Inches(0.5), Inches(1.6), Inches(8), Inches(0.4),
                "ГОРОДСКОЙ СТОЛ ЗАПИСИ · САНКТ-ПЕТЕРБУРГ", size=12, bold=True, color=SIGNAL)
    add_textbox(s, Inches(0.5), Inches(2.2), Inches(10), Inches(1.4),
                "КРУЖИМ", size=72, bold=True, color=WHITE)
    add_textbox(s, Inches(0.5), Inches(3.7), Inches(9), Inches(1.0),
                "Кружки СПб — заявка уходит,\nорганизатор звонит", size=26, bold=True, color=MIST)
    add_textbox(s, Inches(0.5), Inches(5.1), Inches(8.5), Inches(1.0),
                "Позиционирование бренда: территория, сильные стороны,\nпреимущества и почему родители выбирают КРУЖИМ.", size=15, color=SOFT)
    add_textbox(s, Inches(0.5), Inches(6.3), Inches(6), Inches(0.3),
                "кружим.рф", size=13, color=TEAL)
    footer(s, 1, dark=True)

    # --- 2 Territory ---
    s = blank_slide(prs)
    topbar(s, "Территория бренда")
    add_textbox(s, Inches(0.5), Inches(0.85), Inches(12), Inches(0.9),
                "Не каталог. Сервис, где заявка\nзаканчивается звонком.", size=28, bold=True, color=INK)

    # formula card
    add_round_rect(s, Inches(0.5), Inches(2.0), Inches(8.2), Inches(4.6), WHITE)
    lines = [
        ("ДЛЯ", "родителей в Санкт-Петербурге"),
        ("КОТОРЫЕ", "устали искать кружок по чатам, сайтам и «навигаторам»"),
        ("КРУЖИМ", "городской стол записи: возраст + район → короткий список → звонок"),
        ("ПОТОМУ ЧТО", "муниципальные и частные — в одном каталоге, заявка уходит партнёру"),
        ("В ОТЛИЧИЕ ОТ", "Навигатора ДО / Familypass / чата — бюрократия, слот или тишина"),
    ]
    y = 2.2
    for k, v in lines:
        add_textbox(s, Inches(0.75), Inches(y), Inches(1.7), Inches(0.35), k, size=11, bold=True, color=TEAL_DEEP)
        add_textbox(s, Inches(2.5), Inches(y), Inches(5.9), Inches(0.55), v, size=13, color=INK_MID)
        y += 0.65
    add_textbox(s, Inches(0.75), Inches(5.85), Inches(7.5), Inches(0.45),
                "«Кружки СПб — заявка уходит, организатор звонит.»", size=14, bold=True, color=SIGNAL)

    # not for
    add_round_rect(s, Inches(9.0), Inches(2.0), Inches(3.8), Inches(4.6), WHITE)
    add_textbox(s, Inches(9.25), Inches(2.2), Inches(3.3), Inches(0.35),
                "НЕ ДЛЯ КОГО", size=12, bold=True, color=TEAL)
    nots = [
        ("01", "Не маркетплейс разовых слотов"),
        ("02", "Не гос. сертификатный портал"),
        ("03", "Не всероссийский «всё для детей»"),
    ]
    y = 2.8
    for n, t in nots:
        add_textbox(s, Inches(9.25), Inches(y), Inches(0.5), Inches(0.35), n, size=14, bold=True, color=SIGNAL)
        add_textbox(s, Inches(9.85), Inches(y), Inches(2.7), Inches(0.7), t, size=13, color=INK_MID)
        y += 1.0
    footer(s, 2)

    # --- 3 Strengths ---
    s = blank_slide(prs)
    topbar(s, "Сильные стороны")
    add_textbox(s, Inches(0.5), Inches(0.85), Inches(12), Inches(0.6),
                "То, что уже есть в продукте — не обещания «на бумаге»", size=26, bold=True, color=INK)

    strengths = [
        ("01", "Петля доверия:\nзаявка → звонок", "Родитель оставляет заявку, партнёр перезванивает."),
        ("02", "Муниципальное\nи частное", "Бесплатные гос. кружки и платные школы рядом."),
        ("03", "Короткий список,\nне лента", "Возраст + район сразу сужают выбор."),
        ("04", "Модерация\nпартнёров", "Pending → active: не объявления с улицы."),
        ("05", "Глубина СПб,\nне ширина страны", "Районы, карта, локальные организаторы."),
    ]
    for i, (n, title, desc) in enumerate(strengths):
        x = 0.5 + i * 2.5
        add_round_rect(s, Inches(x), Inches(1.8), Inches(2.35), Inches(3.6), WHITE)
        add_textbox(s, Inches(x + 0.15), Inches(2.0), Inches(2.0), Inches(0.4), n, size=18, bold=True, color=TEAL)
        add_textbox(s, Inches(x + 0.15), Inches(2.55), Inches(2.05), Inches(1.1), title, size=14, bold=True, color=INK)
        add_textbox(s, Inches(x + 0.15), Inches(3.9), Inches(2.05), Inches(1.2), desc, size=12, color=MUTED)

    add_round_rect(s, Inches(0.5), Inches(5.65), Inches(12.3), Inches(1.1), RGBColor(0xFF, 0xEB, 0xE3))
    add_textbox(s, Inches(0.7), Inches(5.75), Inches(12), Inches(0.3),
                "КЛЮЧЕВОЙ ВРАГ БРЕНДА", size=11, bold=True, color=SIGNAL)
    add_textbox(s, Inches(0.7), Inches(6.1), Inches(11.8), Inches(0.45),
                "Пустые заявки, мёртвые объявления и чаты без подтверждения записи. Против хаоса поиска — за живой ответ.",
                size=13, color=INK_MID)
    footer(s, 3)

    # --- 4 Parent benefits ---
    s = blank_slide(prs)
    topbar(s, "Преимущества · родители")
    add_textbox(s, Inches(0.5), Inches(0.85), Inches(12), Inches(0.5),
                "Преимущества для родителей", size=28, bold=True, color=INK)
    benefits = [
        ("01", "Экономия вечеров", "Один заход вместо десяти вкладок, чатов мам и «уточните у администратора»."),
        ("02", "Понятный следующий шаг", "Не «сохранить в избранное», а заявка с ожидаемым звонком организатора."),
        ("03", "Честные метки", "Муниципальный / платный, возраст, район, цена или «Бесплатно»."),
        ("04", "Доверие к каталогу", "Партнёры проходят регистрацию и модерацию — не анонимная доска."),
    ]
    for i, (n, title, desc) in enumerate(benefits):
        x = 0.5 + (i % 4) * 3.15
        add_round_rect(s, Inches(x), Inches(1.8), Inches(3.0), Inches(4.2), WHITE)
        add_textbox(s, Inches(x + 0.2), Inches(2.1), Inches(2.5), Inches(0.4), n, size=16, bold=True, color=TEAL)
        add_textbox(s, Inches(x + 0.2), Inches(2.7), Inches(2.6), Inches(1.0), title, size=16, bold=True, color=INK)
        add_textbox(s, Inches(x + 0.2), Inches(3.9), Inches(2.6), Inches(1.6), desc, size=13, color=MUTED)
    footer(s, 4)

    # --- 5 Partners + pillars ---
    s = blank_slide(prs)
    topbar(s, "Партнёры · опоры сообщений")
    add_textbox(s, Inches(0.5), Inches(0.85), Inches(12), Inches(0.5),
                "Для организаторов и три опоры бренда", size=26, bold=True, color=INK)

    add_round_rect(s, Inches(0.5), Inches(1.7), Inches(6.0), Inches(4.8), WHITE)
    add_textbox(s, Inches(0.75), Inches(1.95), Inches(5.5), Inches(0.35),
                "ПРЕИМУЩЕСТВА ДЛЯ ПАРТНЁРОВ", size=12, bold=True, color=TEAL)
    add_textbox(s, Inches(0.75), Inches(2.6), Inches(5.5), Inches(0.4),
                "Заявки от родителей рядом", size=18, bold=True, color=INK)
    add_textbox(s, Inches(0.75), Inches(3.1), Inches(5.5), Inches(0.8),
                "Не «место в каталоге ради галочки», а лиды из целевого района и возраста.", size=14, color=MUTED)
    add_textbox(s, Inches(0.75), Inches(4.2), Inches(5.5), Inches(0.4),
                "Кабинет и статусы", size=18, bold=True, color=INK)
    add_textbox(s, Inches(0.75), Inches(4.7), Inches(5.5), Inches(0.8),
                "Управление кружками, модерация, путь от регистрации до публикации.", size=14, color=MUTED)

    add_round_rect(s, Inches(6.8), Inches(1.7), Inches(6.0), Inches(4.8), DARK_BG)
    add_textbox(s, Inches(7.05), Inches(1.95), Inches(5.5), Inches(0.35),
                "ТРИ ОПОРЫ СООБЩЕНИЙ", size=12, bold=True, color=SIGNAL)
    pillars = [
        ("Короткий список", "Сначала возраст и район — потом варианты"),
        ("Живой ответ", "Заявка заканчивается звонком, не тишиной"),
        ("Всё честное рядом", "Муниципальное и частное — с явной меткой"),
    ]
    y = 2.6
    for title, desc in pillars:
        add_textbox(s, Inches(7.05), Inches(y), Inches(5.5), Inches(0.35), title, size=16, bold=True, color=WHITE)
        add_textbox(s, Inches(7.05), Inches(y + 0.4), Inches(5.5), Inches(0.45), desc, size=13, color=SOFT)
        y += 1.15
    footer(s, 5)

    # --- 6 Comparison ---
    s = blank_slide(prs)
    topbar(s, "Почему КРУЖИМ")
    add_textbox(s, Inches(0.5), Inches(0.8), Inches(12), Inches(0.45),
                "Почему КРУЖИМ, а не кто-то ещё", size=26, bold=True, color=INK)
    add_textbox(s, Inches(0.5), Inches(1.3), Inches(12), Inches(0.35),
                "Главный вопрос родителя: «Мне ответят и запишут — или снова в пустоту?»", size=14, color=MUTED)

    rows = [
        ("Альтернатива", "Их обещание", "Слабое место", "КРУЖИМ"),
        ("Навигатор ДО", "Госреестр и сертификаты", "Бюрократия, не про быструю запись", "Живая заявка + звонок"),
        ("Familypass", "«Все секции», покупка слотов", "Ширина важнее подтверждения", "Глубина СПб + ответ"),
        ("Чаты / форумы", "«Живые советы»", "Хаос, без статуса заявки", "Структура + статус"),
        ("Google / Avito", "Найти объявление", "Мёртвые страницы", "Модерация партнёров"),
        ("Ничего не делать", "Подождать / спросить в школе", "Упущенный сезон", "Старт за вечер"),
    ]
    col_w = [2.4, 3.4, 3.6, 2.9]
    y = 1.8
    for ri, row in enumerate(rows):
        x = 0.5
        bg = RGBColor(0x12, 0x30, 0x40) if ri == 0 else (WHITE if ri % 2 else FOG)
        h = 0.55 if ri == 0 else 0.62
        for ci, (cell, w) in enumerate(zip(row, col_w)):
            add_rect(s, Inches(x), Inches(y), Inches(w), Inches(h), bg)
            c = WHITE if ri == 0 else (TEAL_DEEP if ci == 3 else INK_MID)
            bold = ri == 0 or ci in (0, 3)
            add_textbox(s, Inches(x + 0.08), Inches(y + 0.12), Inches(w - 0.16), Inches(h - 0.1),
                        cell, size=11 if ri else 10, bold=bold, color=c)
            x += w
        y += h

    add_textbox(s, Inches(0.5), Inches(6.2), Inches(12), Inches(0.45),
                "Другие собирают занятия. КРУЖИМ доводит до разговора с организатором.",
                size=16, bold=True, color=INK)
    footer(s, 6)

    # --- 7 Voice ---
    s = blank_slide(prs)
    topbar(s, "Голос бренда")
    add_textbox(s, Inches(0.5), Inches(0.85), Inches(12), Inches(0.5),
                "Как говорим — и как не говорим", size=28, bold=True, color=INK)

    add_round_rect(s, Inches(0.5), Inches(1.7), Inches(6.0), Inches(4.8), WHITE)
    add_textbox(s, Inches(0.75), Inches(1.95), Inches(5.5), Inches(0.35),
                "КАК ГОВОРИМ", size=13, bold=True, color=TEAL_DEEP)
    dos = [
        "Спокойно и конкретно",
        "Глаголами процесса: выбрали → оставили → перезвонили",
        "Уважаем время родителя",
        "Без сюсюканья и крика",
        "Городской тон СПб",
    ]
    y = 2.55
    for t in dos:
        add_textbox(s, Inches(0.75), Inches(y), Inches(5.5), Inches(0.45), f"•  {t}", size=14, color=INK_MID)
        y += 0.55

    add_round_rect(s, Inches(6.8), Inches(1.7), Inches(6.0), Inches(4.8), WHITE)
    add_textbox(s, Inches(7.05), Inches(1.95), Inches(5.5), Inches(0.35),
                "КАК НЕ ГОВОРИМ", size=13, bold=True, color=SIGNAL)
    donts = [
        "«Лучший выбор», «#1», «экосистема»",
        "«Инновации», пустая «забота»",
        "«Всё для детей» без фокуса",
        "Гонка числом занятий с агрегаторами",
        "Притворство гос. навигатором",
    ]
    y = 2.55
    for t in donts:
        add_textbox(s, Inches(7.05), Inches(y), Inches(5.5), Inches(0.45), f"•  {t}", size=14, color=INK_MID)
        y += 0.55
    footer(s, 7)

    # --- 8 Hero + proof ---
    s = blank_slide(prs)
    topbar(s, "Hero · Proof roadmap")
    add_textbox(s, Inches(0.5), Inches(0.85), Inches(12), Inches(0.45),
                "Сообщение = продукт", size=28, bold=True, color=INK)

    add_round_rect(s, Inches(0.5), Inches(1.6), Inches(6.2), Inches(4.9), DARK_BG)
    add_textbox(s, Inches(0.9), Inches(2.2), Inches(5.4), Inches(0.6),
                "КРУЖИМ", size=32, bold=True, color=WHITE)
    add_textbox(s, Inches(0.9), Inches(3.0), Inches(5.4), Inches(0.9),
                "Заявка в кружок — и звонок организатора", size=20, bold=True, color=MIST)
    add_textbox(s, Inches(0.9), Inches(4.1), Inches(5.4), Inches(0.7),
                "Возраст и район → муниципальные и частные занятия в СПб", size=14, color=SOFT)
    cta = add_round_rect(s, Inches(0.9), Inches(5.2), Inches(2.8), Inches(0.55), SIGNAL)
    add_textbox(s, Inches(0.9), Inches(5.28), Inches(2.8), Inches(0.4),
                "Подобрать кружок", size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

    proofs = [
        ("1", "Метрика «% заявок с ответом за 24ч»", "Сделать «живой ответ» измеримым"),
        ("2", "SLA / пауза партнёров без звонка", "Защитить доверие"),
        ("3", "Hero и партнёрский лендинг на spine", "Сообщение = продукт"),
        ("4", "Метки муниципальный / платный", "Честность каталога"),
    ]
    y = 1.6
    for n, title, desc in proofs:
        add_round_rect(s, Inches(7.0), Inches(y), Inches(5.8), Inches(1.1), WHITE)
        add_textbox(s, Inches(7.2), Inches(y + 0.2), Inches(0.5), Inches(0.4), n, size=16, bold=True, color=TEAL)
        add_textbox(s, Inches(7.8), Inches(y + 0.15), Inches(4.7), Inches(0.35), title, size=13, bold=True, color=INK)
        add_textbox(s, Inches(7.8), Inches(y + 0.55), Inches(4.7), Inches(0.35), desc, size=12, color=MUTED)
        y += 1.2
    footer(s, 8)

    # --- 9 Close ---
    s = blank_slide(prs, dark=True)
    topbar(s, "Следующий шаг", dark=True)
    add_textbox(s, Inches(0.5), Inches(1.8), Inches(10), Inches(0.4),
                "КРУЖИМ.РФ · САНКТ-ПЕТЕРБУРГ", size=12, bold=True, color=SIGNAL)
    add_textbox(s, Inches(0.5), Inches(2.4), Inches(12), Inches(1.4),
                "Заявка уходит.\nОрганизатор звонит.", size=40, bold=True, color=WHITE)
    add_textbox(s, Inches(0.5), Inches(4.2), Inches(10), Inches(0.8),
                "Городской стол записи в детские занятия — короткий список\nпо возрасту и району, живой ответ партнёра.", size=16, color=SOFT)
    add_round_rect(s, Inches(0.5), Inches(5.4), Inches(3.2), Inches(0.6), SIGNAL)
    add_textbox(s, Inches(0.5), Inches(5.5), Inches(3.2), Inches(0.4),
                "Подобрать кружок →", size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    footer(s, 9, dark=True)

    out = "/workspace/presentation/КРУЖИМ_Brand_Positioning_2026.pptx"
    prs.save(out)
    print(f"Saved: {out}")
    return out


if __name__ == "__main__":
    build()
