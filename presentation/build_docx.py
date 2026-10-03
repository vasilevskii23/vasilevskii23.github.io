#!/usr/bin/env python3
"""КРУЖИМ Brand Positioning 2026 — DOCX in brand style from source PDF."""

from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches, Twips, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn, nsmap
from docx.oxml import OxmlElement

ROOT = Path(__file__).resolve().parent
LOGO = ROOT / "assets" / "logo.png"
OUT = ROOT / "KRUZHIM_Brand_Positioning_2026.docx"

# Brand colors
NAVY = RGBColor(0x0D, 0x1F, 0x5C)
BLUE = RGBColor(0x1C, 0x42, 0x99)
BLUE_MID = RGBColor(0x2A, 0x5B, 0xC4)
SOFT = RGBColor(0x79, 0x95, 0xCC)
MUTED = RGBColor(0x4A, 0x5F, 0x8F)
ORANGE = RGBColor(0xE8, 0x7A, 0x2E)
AMBER = RGBColor(0xD4, 0x92, 0x3F)
GREEN = RGBColor(0x16, 0xA2, 0x4A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BG_HEX = "F0F5FF"
CARD_HEX = "EAF0FF"
NAVY_HEX = "0D1F5C"
BLUE_HEX = "1C4299"
ORANGE_HEX = "E87A2E"
GREEN_HEX = "16A24A"
SOFT_HEX = "7995CC"


def set_run_font(run, size=11, bold=False, color=NAVY, name="Arial"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def shade_cell(cell, hex_color):
    tc = cell._tePr if hasattr(cell, "_tePr") else cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_borders(cell, color="D0DAF0", sz="4"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:color"), color)
        borders.append(el)
    tcPr.append(borders)


def clear_cell_paragraphs(cell):
    for p in cell.paragraphs:
        p.clear()


def add_para(doc_or_cell, text, size=11, bold=False, color=NAVY, space_before=0, space_after=6, align=WD_ALIGN_PARAGRAPH.LEFT):
    if hasattr(doc_or_cell, "add_paragraph"):
        p = doc_or_cell.add_paragraph()
    else:
        # cell
        p = doc_or_cell.paragraphs[0] if not doc_or_cell.paragraphs[0].text else doc_or_cell.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.25
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return p


def add_mixed_para(container, parts, space_before=0, space_after=6, align=WD_ALIGN_PARAGRAPH.LEFT):
    """parts: list of (text, size, bold, color)"""
    if hasattr(container, "add_paragraph"):
        p = container.add_paragraph()
    else:
        p = container.paragraphs[0] if not container.paragraphs[0].text else container.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.3
    for text, size, bold, color in parts:
        run = p.add_run(text)
        set_run_font(run, size=size, bold=bold, color=color)
    return p


def section_heading(doc, title, eyebrow=None):
    if eyebrow:
        add_para(doc, eyebrow.upper(), size=10, bold=True, color=SOFT, space_before=18, space_after=4)
    add_para(doc, title, size=20, bold=True, color=NAVY, space_before=0, space_after=10)


def add_hr(doc, color=SOFT_HEX):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def card_table(doc, rows_content, col_count, fill=CARD_HEX):
    """Create a grid of cards. rows_content is list of lists of (title, body)."""
    table = doc.add_table(rows=len(rows_content), cols=col_count)
    table.autofit = True
    for ri, row in enumerate(rows_content):
        for ci in range(col_count):
            cell = table.rows[ri].cells[ci]
            shade_cell(cell, fill)
            set_cell_borders(cell, "D8E2F5")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            # padding via margins
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            mar = OxmlElement("w:tcMar")
            for edge, val in (("top", "80"), ("left", "100"), ("bottom", "80"), ("right", "100")):
                el = OxmlElement(f"w:{edge}")
                el.set(qn("w:w"), val)
                el.set(qn("w:type"), "dxa")
                mar.append(el)
            tcPr.append(mar)
            if ci < len(row):
                title, body = row[ci]
                clear_cell_paragraphs(cell)
                add_para(cell, title, size=12, bold=True, color=NAVY, space_after=4)
                add_para(cell, body, size=10, bold=False, color=MUTED, space_after=2)
            else:
                clear_cell_paragraphs(cell)
                cell.paragraphs[0].text = ""
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return table


def build():
    doc = Document()

    # Page setup
    for section in doc.sections:
        section.top_margin = Cm(1.8)
        section.bottom_margin = Cm(1.8)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)

    # --- COVER / HEADER ---
    header_table = doc.add_table(rows=1, cols=2)
    header_table.autofit = True
    c0, c1 = header_table.rows[0].cells
    # logo + brand
    p = c0.paragraphs[0]
    if LOGO.exists():
        run = p.add_run()
        run.add_picture(str(LOGO), width=Cm(0.9))
        p.add_run("  ")
    run = p.add_run("КРУЖИМ")
    set_run_font(run, size=16, bold=True, color=NAVY)
    p1 = c1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p1.add_run("BRAND POSITIONING · 2026")
    set_run_font(run, size=10, bold=True, color=SOFT)

    add_hr(doc, BLUE_HEX)

    # Cover block as shaded table
    cover = doc.add_table(rows=1, cols=1)
    cell = cover.rows[0].cells[0]
    shade_cell(cell, BLUE_HEX)
    set_cell_borders(cell, BLUE_HEX)
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for edge, val in (("top", "160"), ("left", "160"), ("bottom", "160"), ("right", "160")):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), val)
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tcPr.append(mar)
    clear_cell_paragraphs(cell)
    add_para(cell, "BRAND POSITIONING · 2026", size=10, bold=True, color=SOFT, space_after=8)
    add_para(cell, "Кружки СПб — заявка уходит,", size=22, bold=True, color=WHITE, space_after=2)
    add_para(cell, "организатор звонит", size=22, bold=True, color=ORANGE, space_after=12)
    add_para(
        cell,
        "Позиционирование бренда: территория, сильные стороны, преимущества и почему родители выбирают КРУЖИМ, а не «кого-то ещё».",
        size=11,
        color=SOFT,
        space_after=10,
    )
    add_para(cell, "Городской стол записи в детские занятия  ·  кружим.рф · Санкт-Петербург", size=10, bold=True, color=WHITE, space_after=2)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # --- 1 TERRITORY ---
    section_heading(doc, "Территория бренда", "Позиционирование")
    add_para(
        doc,
        "Живое подтверждение — не «ещё один каталог», а сервис, где заявка заканчивается звонком организатора.",
        size=11,
        color=MUTED,
        space_after=10,
    )

    pos = doc.add_table(rows=1, cols=1)
    cell = pos.rows[0].cells[0]
    shade_cell(cell, NAVY_HEX)
    set_cell_borders(cell, NAVY_HEX)
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for edge, val in (("top", "120"), ("left", "140"), ("bottom", "120"), ("right", "140")):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), val)
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tcPr.append(mar)
    clear_cell_paragraphs(cell)

    formula = [
        ("Для ", "родителей в Санкт-Петербурге"),
        ("которые ", "устали искать кружок по чатам, сайтам и «навигаторам» и не понимать, живая ли запись"),
        ("КРУЖИМ — ", "городской стол записи в детские занятия, который из возраста и района собирает короткий список и доводит до звонка организатора"),
        ("потому что ", "муниципальные и частные кружки — в одном каталоге, а заявка уходит реальному партнёру"),
        ("в отличие от ", "Навигатора ДО / Familypass / районного чата, где либо бюрократия, либо покупка слота, либо тишина"),
    ]
    for k, v in formula:
        add_mixed_para(
            cell,
            [(k, 11, True, AMBER), (v, 11, False, WHITE)],
            space_after=6,
        )
    add_para(cell, "«Кружки СПб — заявка уходит, организатор звонит.»", size=12, bold=True, color=WHITE, space_before=8, space_after=2)

    add_para(doc, "НЕ ДЛЯ КОГО", size=10, bold=True, color=SOFT, space_before=12, space_after=6)
    card_table(
        doc,
        [[
            ("01  Не маркетплейс разовых слотов", "Не продаём разовые слоты как товар."),
            ("02  Не гос. сертификатный портал", "Не подменяем Навигатор ДО."),
            ("03  Не всероссийский «всё для детей»", "Фокус — Санкт-Петербург."),
        ]],
        3,
        fill=CARD_HEX,
    )

    # --- 2 STRENGTHS ---
    section_heading(doc, "Сильные стороны продукта", "Сильные стороны")
    add_para(
        doc,
        "То, что уже есть в продукте и защищает позиционирование — не обещания «на бумаге».",
        size=11,
        color=MUTED,
        space_after=8,
    )
    strengths = [
        ("1. Петля доверия: заявка → звонок", "Родитель оставляет заявку, партнёр получает лид и перезванивает. Это отличает КРУЖИМ от мёртвых форм и объявлений."),
        ("2. Муниципальное и частное в одном окне", "Бесплатные гос. кружки с явной меткой и платные школы рядом — без переключения между «двумя мирами»."),
        ("3. Короткий список, не бесконечная лента", "Возраст + район (+ карта) сразу сужают выбор. Меньше шума — быстрее решение."),
        ("4. Модерация партнёров", "Кабинет партнёра и статус pending → active: в каталог попадают не «любые объявления с улицы»."),
        ("5. Глубина СПб, а не ширина страны", "Фокус на одном городе: районы, карта, локальные организаторы — вместо гонки «7000 занятий в 16 городах»."),
    ]
    for title, body in strengths:
        add_para(doc, title, size=12, bold=True, color=NAVY, space_before=4, space_after=2)
        add_para(doc, body, size=10, color=MUTED, space_after=6)

    enemy = doc.add_table(rows=1, cols=1)
    cell = enemy.rows[0].cells[0]
    shade_cell(cell, "FFF0E4")
    set_cell_borders(cell, ORANGE_HEX)
    clear_cell_paragraphs(cell)
    add_para(cell, "Ключевой враг бренда", size=11, bold=True, color=ORANGE, space_after=4)
    add_para(
        cell,
        "Пустые заявки, мёртвые объявления и чаты, где никто не подтверждает запись. Мы против хаоса поиска — за живой ответ.",
        size=11,
        color=NAVY,
        space_after=2,
    )

    # --- 3 BENEFITS ---
    section_heading(doc, "Преимущества для родителей", "Преимущества")
    card_table(
        doc,
        [
            [
                ("Экономия вечеров", "Один заход вместо десяти вкладок, чатов мам и «уточните у администратора»."),
                ("Понятный следующий шаг", "Не «сохранить в избранное», а заявка с ожидаемым звонком организатора."),
            ],
            [
                ("Честные метки", "Видно: муниципальный / платный, возраст, район, цена или «Бесплатно»."),
                ("Доверие к каталогу", "Партнёры проходят регистрацию и модерацию — не анонимная доска объявлений."),
            ],
        ],
        2,
    )

    add_para(doc, "ПРЕИМУЩЕСТВА ДЛЯ ПАРТНЁРОВ (ОРГАНИЗАТОРОВ)", size=10, bold=True, color=BLUE, space_before=4, space_after=6)
    card_table(
        doc,
        [[
            ("Заявки от родителей рядом", "Не «место в каталоге ради галочки», а лиды из целевого района и возраста."),
            ("Кабинет и статусы", "Управление кружками, модерация, понятный путь от регистрации до публикации."),
        ]],
        2,
    )

    add_para(doc, "ТРИ ОПОРЫ СООБЩЕНИЙ", size=10, bold=True, color=BLUE, space_before=4, space_after=6)
    pillars = doc.add_table(rows=1, cols=3)
    for i, (title, body) in enumerate([
        ("Короткий список", "Сначала возраст и район — потом варианты"),
        ("Живой ответ", "Заявка заканчивается звонком, не тишиной"),
        ("Всё честное рядом", "Муниципальное и частное — с явной меткой"),
    ]):
        cell = pillars.rows[0].cells[i]
        shade_cell(cell, BLUE_HEX)
        set_cell_borders(cell, BLUE_HEX)
        clear_cell_paragraphs(cell)
        add_para(cell, title, size=12, bold=True, color=WHITE, space_after=4)
        add_para(cell, body, size=10, color=SOFT, space_after=2)

    # --- 4 WHY ---
    section_heading(doc, "Почему КРУЖИМ, а не кто-то ещё", "Почему КРУЖИМ")
    add_para(
        doc,
        "Сравнение по главному вопросу родителя: «Мне ответят и запишут — или снова в пустоту?»",
        size=11,
        color=MUTED,
        space_after=8,
    )

    rows = [
        ("Альтернатива", "Их обещание", "Слабое место", "КРУЖИМ"),
        ("Навигатор ДО", "Госреестр программ и сертификаты", "Бюрократия, UX, не про «быстро записаться»", "Живая заявка + звонок"),
        ("Familypass и аналоги", "«Все секции», покупка слотов", "Ширина каталога важнее подтверждения записи", "Глубина СПб + ответ"),
        ("Чаты / мама-форумы", "«Живые советы»", "Хаос, устаревшее, без статуса заявки", "Структура + статус"),
        ("Google / Avito", "Найти объявление", "Мёртвые страницы, нет модерации", "Модерация партнёров"),
        ("«Ничего не делать»", "Подождать / спросить у школы", "Месяцы без кружка, упущенный сезон", "Старт за вечер"),
    ]
    table = doc.add_table(rows=len(rows), cols=4)
    table.style = "Table Grid"
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.rows[ri].cells[ci]
            clear_cell_paragraphs(cell)
            if ri == 0:
                shade_cell(cell, BLUE_HEX)
                add_para(cell, val, size=9, bold=True, color=WHITE, space_after=2)
            else:
                shade_cell(cell, "FFFFFF" if ri % 2 else CARD_HEX)
                if ci == 3:
                    color = ORANGE if ri == 5 else GREEN
                    add_para(cell, val, size=9, bold=True, color=color, space_after=2)
                elif ci == 0:
                    add_para(cell, val, size=9, bold=True, color=NAVY, space_after=2)
                else:
                    add_para(cell, val, size=9, color=MUTED, space_after=2)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    one = doc.add_table(rows=1, cols=1)
    cell = one.rows[0].cells[0]
    shade_cell(cell, "FFF0E4")
    set_cell_borders(cell, ORANGE_HEX)
    clear_cell_paragraphs(cell)
    add_para(cell, "Одной фразой", size=10, bold=True, color=ORANGE, space_after=4)
    add_para(
        cell,
        "Другие собирают занятия. КРУЖИМ доводит до разговора с организатором.",
        size=12,
        bold=True,
        color=NAVY,
        space_after=2,
    )

    # --- 5 VOICE ---
    section_heading(doc, "Голос бренда и proof roadmap", "Голос и следующий шаг")
    voice = doc.add_table(rows=1, cols=2)
    left, right = voice.rows[0].cells
    shade_cell(left, CARD_HEX)
    shade_cell(right, CARD_HEX)
    set_cell_borders(left, "D8E2F5")
    set_cell_borders(right, "D8E2F5")
    clear_cell_paragraphs(left)
    clear_cell_paragraphs(right)
    add_para(left, "Как говорим", size=12, bold=True, color=NAVY, space_after=6)
    for t in [
        "Спокойно и конкретно",
        "Глаголами процесса: выбрали → оставили → перезвонили",
        "Уважаем время родителя",
        "Без сюсюканья и крика",
        "Городской тон СПб",
    ]:
        add_para(left, f"•  {t}", size=10, color=MUTED, space_after=3)
    add_para(right, "Как не говорим", size=12, bold=True, color=NAVY, space_after=6)
    for t in [
        "«Лучший выбор», «#1», «экосистема»",
        "«Инновации», пустая «забота»",
        "«Всё для детей» без фокуса",
        "Гонка числом занятий с агрегаторами",
        "Притворство гос. навигатором",
    ]:
        add_para(right, f"•  {t}", size=10, color=MUTED, space_after=3)

    add_para(doc, "HERO (ЧЕРНОВИК ПОД ПОЗИЦИОНИРОВАНИЕ)", size=10, bold=True, color=SOFT, space_before=12, space_after=6)
    hero = doc.add_table(rows=1, cols=1)
    cell = hero.rows[0].cells[0]
    shade_cell(cell, BLUE_HEX)
    set_cell_borders(cell, BLUE_HEX)
    clear_cell_paragraphs(cell)
    add_para(cell, "КРУЖИМ", size=14, bold=True, color=WHITE, space_after=4)
    add_para(cell, "Заявка в кружок — и звонок организатора", size=14, bold=True, color=WHITE, space_after=4)
    add_para(cell, "Возраст и район → муниципальные и частные занятия в СПб", size=11, color=SOFT, space_after=8)
    add_para(cell, "CTA: Подобрать кружок", size=11, bold=True, color=ORANGE, space_after=2)

    add_para(doc, "ЧТО УСИЛИТЬ В ПРОДУКТЕ (ЧТОБЫ БРЕНД НЕ ВРАЛ)", size=10, bold=True, color=BLUE, space_before=12, space_after=6)
    proofs = [
        ("#", "PROOF", "ЗАЧЕМ"),
        ("1", "Метрика «% заявок с ответом за 24ч»", "Сделать «живой ответ» измеримым"),
        ("2", "SLA / пауза партнёров, которые не звонят", "Защитить доверие"),
        ("3", "Hero и партнёрский лендинг на spine", "Сообщение = продукт"),
        ("4", "Показ муниципальный / платный всегда явно", "Честность каталога"),
    ]
    pt = doc.add_table(rows=len(proofs), cols=3)
    for ri, row in enumerate(proofs):
        for ci, val in enumerate(row):
            cell = pt.rows[ri].cells[ci]
            clear_cell_paragraphs(cell)
            if ri == 0:
                shade_cell(cell, CARD_HEX)
                add_para(cell, val, size=9, bold=True, color=SOFT, space_after=2)
            else:
                shade_cell(cell, "FFFFFF")
                add_para(cell, val, size=10, bold=(ci == 0), color=NAVY if ci < 2 else MUTED, space_after=2)

    # Footer
    add_hr(doc, SOFT_HEX)
    add_para(doc, "КРУЖИМ · Brand book · кружим.рф · Санкт-Петербург", size=9, color=SOFT, align=WD_ALIGN_PARAGRAPH.CENTER)

    doc.save(OUT)
    print(f"Saved: {OUT}")
    return OUT


if __name__ == "__main__":
    build()
